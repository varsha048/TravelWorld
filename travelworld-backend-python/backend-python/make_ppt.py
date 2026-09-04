import collections.abc
import collections
# Monkey patch for python-pptx compatibility with python 3.10+
collections.Iterator = collections.abc.Iterator
collections.Container = collections.abc.Container
collections.Sequence = collections.abc.Sequence
collections.Set = collections.abc.Set
collections.Mapping = collections.abc.Mapping
collections.MutableSet = collections.abc.MutableSet
collections.MutableMapping = collections.abc.MutableMapping
collections.Iterable = collections.abc.Iterable

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Slide 1: Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "TravelWorld"
subtitle.text = "An AI-Powered Tour Booking & Recommendation Platform\n\nAdvanced Modules Overview"

# Slide 2: Introduction
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Introduction & Problem Statement"
tf = body_shape.text_frame
tf.text = "The Modern Travel Planning Problem:"
p = tf.add_paragraph()
p.text = "Fragmented tools: Users switch between booking, maps, and weather sites."
p.level = 1
p = tf.add_paragraph()
p.text = "Time-consuming: Building an itinerary requires hours of manual research."
p.level = 1
p = tf.add_paragraph()
p.text = "TravelWorld Solution:"
p = tf.add_paragraph()
p.text = "A centralized platform providing real-time data, AI-driven itineraries, and instant PDF ticketing."
p.level = 1

# Slide 3: Scope & Objectives
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Scope & Objectives"
tf = body_shape.text_frame
tf.text = "Core Objectives:"
p = tf.add_paragraph()
p.text = "Develop a cohesive UI with real-time environmental context."
p.level = 1
p = tf.add_paragraph()
p.text = "Automate travel planning using Generative AI (Gemini)."
p.level = 1
p = tf.add_paragraph()
p.text = "Streamline post-booking with automated PDF E-Tickets."
p.level = 1
p = tf.add_paragraph()
p.text = "Scope Boundary:"
p = tf.add_paragraph()
p.text = "Focuses on advanced integrations (AI, Maps, Analytics, PDF Gen)."
p.level = 1
p = tf.add_paragraph()
p.text = "Excludes standard login and basic tour browsing modules."
p.level = 1

# Slide 4: System Architecture
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "System Architecture"
tf = body_shape.text_frame
tf.text = "Hybrid 3-Tier & MVC Architecture"
p = tf.add_paragraph()
p.text = "Presentation Layer (Client): HTML5, CSS3, JavaScript."
p.level = 1
p = tf.add_paragraph()
p.text = "Business Logic (Server): Python Flask handling AI proxies and ticketing."
p.level = 1
p = tf.add_paragraph()
p.text = "Data Layer: SQLite database storing Bookings and Transactions."
p.level = 1
p = tf.add_paragraph()
p.text = "Third-Party APIs: Gemini AI, Open-Meteo, Leaflet.js."
p.level = 1

# Slide 5: Advanced Features (Part 1)
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Key Modules: Context & AI"
tf = body_shape.text_frame
tf.text = "1. Live Environmental Data"
p = tf.add_paragraph()
p.text = "Interactive Maps via Leaflet.js and OpenStreetMap."
p.level = 1
p = tf.add_paragraph()
p.text = "Live 3-day weather forecasting via Open-Meteo API."
p.level = 1
p = tf.add_paragraph()
p.text = "2. AI Trip Generator"
p = tf.add_paragraph()
p.text = "Integrates Gemini 2.5 Flash for custom itinerary creation."
p.level = 1
p = tf.add_paragraph()
p.text = "Instantly generates a day-by-day plan based on budget and interests."
p.level = 1

# Slide 6: Advanced Features (Part 2)
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Key Modules: Transactions & Admin"
tf = body_shape.text_frame
tf.text = "3. Booking & E-Ticketing"
p = tf.add_paragraph()
p.text = "Simulated secure checkout modal."
p.level = 1
p = tf.add_paragraph()
p.text = "Automated PDF Ticket generation using ReportLab."
p.level = 1
p = tf.add_paragraph()
p.text = "4. Administrative Analytics"
p = tf.add_paragraph()
p.text = "Protected dashboard for business operators."
p.level = 1
p = tf.add_paragraph()
p.text = "Visual data rendering via Chart.js (Revenue, Booking Volume)."
p.level = 1

# Slide 7: Conclusion
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Conclusion"
tf = body_shape.text_frame
tf.text = "Impact of TravelWorld:"
p = tf.add_paragraph()
p.text = "Effectively centralizes fragmented travel tools into a single platform."
p.level = 1
p = tf.add_paragraph()
p.text = "Drastically reduces planning time using Generative AI."
p.level = 1
p = tf.add_paragraph()
p.text = "Provides robust, visual tools for business administrators."
p.level = 1
p = tf.add_paragraph()
p.text = "Demonstrates modern web design and complex third-party API integration."
p.level = 1

# Save presentation
prs.save('TravelWorld_Presentation.pptx')
print("Presentation generated successfully at TravelWorld_Presentation.pptx")
