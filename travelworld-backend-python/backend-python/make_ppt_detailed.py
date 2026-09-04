import collections.abc
import collections
# Monkey patch for python-pptx compatibility with python 3.10+
collections.Iterator = collections.abc.Iterator
collections.Container = collections.abc.Container
collections.Sequence = collections.abc.Sequence
collections.Set = collections.abc.Set
collections.Mapping = collections.abc.Mapping
collections.MutableSet = collections.abc.MutableSet
collections.MutableMapping = collections.abc.MutableMapping
collections.Iterable = collections.abc.Iterable

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    # A rich, modern dark blue/teal color for the background
    fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900

def format_title(title_shape, text):
    title_shape.text = text
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
        paragraph.font.name = 'Segoe UI'
        paragraph.font.bold = True

def format_body(body_shape, text_lines):
    tf = body_shape.text_frame
    tf.clear()
    for i, line in enumerate(text_lines):
        p = tf.add_paragraph()
        
        # Handle indentations (lines starting with '-')
        if line.startswith("- "):
            p.text = line[2:]
            p.level = 1
        elif line.startswith("-- "):
            p.text = line[3:]
            p.level = 2
        else:
            p.text = line
            p.level = 0
            
        p.font.color.rgb = RGBColor(241, 245, 249) # Slate 100 (Off-white)
        p.font.name = 'Segoe UI'
        if p.level == 0:
            p.font.size = Pt(24)
        else:
            p.font.size = Pt(20)

prs = Presentation()

# --- 1. Title Slide ---
slide = prs.slides.add_slide(prs.slide_layouts[0])
set_slide_background(slide)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "TravelWorld"
for p in title.text_frame.paragraphs:
    p.font.color.rgb = RGBColor(56, 189, 248)
    p.font.bold = True
subtitle.text = "Tour Booking & Recommendation Platform\nFinal Project Presentation"
for p in subtitle.text_frame.paragraphs:
    p.font.color.rgb = RGBColor(241, 245, 249)

# Helper for bullet slides
def add_bullet_slide(title_text, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_slide_background(slide)
    format_title(slide.shapes.title, title_text)
    format_body(slide.placeholders[1], lines)
    return slide

# --- 2. Abstract ---
add_bullet_slide("1. Abstract", [
    "Travel planning is currently a highly fragmented and stressful process.",
    "TravelWorld is a centralized, full-stack web application built to solve this.",
    "It consolidates real-time weather, interactive maps, and AI-driven itineraries.",
    "The platform features automated PDF ticketing and simulated payment checkouts.",
    "Result: A seamless, all-in-one travel e-commerce solution."
])

# --- 3. Introduction ---
add_bullet_slide("2. Introduction", [
    "The Shift to Digital Travel:",
    "- Modern travelers prefer self-service online platforms over traditional agents.",
    "The Information Overload:",
    "- Users must juggle booking engines, mapping tools, and weather apps.",
    "The AI Revolution:",
    "- Leveraging Generative AI (Gemini) allows the system to act as a digital agent.",
    "Project Goal:",
    "- Unify disparate data sources into one intelligent, easy-to-use interface."
])

# --- 4. Problem Statement ---
add_bullet_slide("3. Problem Statement", [
    "Fragmented Planning (Tab-Switching Fatigue)",
    "- Users leave booking sites to research weather and locations.",
    "Time-Consuming Itinerary Creation",
    "- Building a custom schedule takes hours of manual work.",
    "Lack of True Personalization",
    "- Existing sites offer basic filters, not intelligent recommendations.",
    "Inefficient Administrative Tools",
    "- Small operators lack visual, real-time data analytics."
])

# --- 5. Objectives ---
add_bullet_slide("4. Objectives", [
    "Develop a centralized travel hub for browsing and booking.",
    "Integrate Live Environmental Context (Open-Meteo, Leaflet.js).",
    "Automate personalized trip planning using Generative AI.",
    "Streamline post-booking with instant PDF E-Ticket generation.",
    "Provide Administrators with a visual data dashboard (Chart.js)."
])

# --- 6. Existing vs Proposed ---
add_bullet_slide("5. Existing vs. Proposed System", [
    "Existing Systems:",
    "- Fragmented data (Maps and Weather are on separate websites).",
    "- Manual itinerary creation required.",
    "- Cluttered UI with heavy cognitive load.",
    "Proposed System (TravelWorld):",
    "- Unified context (Live Maps & Weather embedded on the tour page).",
    "- Instant Generative AI trip builder.",
    "- Clean, Glassmorphism UI ensuring a premium experience."
])

# --- 7. System Requirements ---
add_bullet_slide("6. System Requirements", [
    "Hardware Requirements:",
    "- Minimum: Intel Core i3, 4GB RAM, 10GB Storage",
    "- Recommended: Intel Core i5, 8GB RAM, SSD Storage",
    "Software Requirements:",
    "- Frontend: HTML5, CSS3, Vanilla JS, Chart.js",
    "- Backend: Python 3, Flask framework, ReportLab",
    "- Database: SQLite (Relational Storage)",
    "- APIs: Gemini 2.5 Flash, Open-Meteo, OpenStreetMap"
])

# --- 8. Modules ---
add_bullet_slide("7. Advanced Modules", [
    "Interactive Mapping & Weather",
    "- Real-time geolocation rendering and 3-day forecasting.",
    "AI Trip Generator",
    "- Dynamic JSON prompt engineering via Gemini for custom travel plans.",
    "Booking & PDF Ticketing",
    "- Secure checkout flow with automated ReportLab PDF rendering.",
    "Administrative Analytics",
    "- Data aggregation and visual charting of total revenue and popularity."
])

# --- 9. System Design (Diagrams Placeholder) ---
design_slides = [
    ("8. System Design: Data Flow Diagram", "[Please Insert Data Flow Diagram (DFD) Screenshot Here]"),
    ("8. System Design: Use Case Diagram", "[Please Insert Use Case Diagram Screenshot Here]"),
    ("8. System Design: Sequence Diagram", "[Please Insert Sequence Diagram Screenshot Here]\n\n(Shows the step-by-step sequence of the AI Generation & Booking process)"),
    ("8. System Design: Entity Relationship (ER)", "[Please Insert ER Diagram Screenshot Here]\n\n(Shows relations between Users, Tours, and Bookings)")
]

for title, placeholder in design_slides:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_slide_background(slide)
    format_title(slide.shapes.title, title)
    
    # Create a nice placeholder box in the middle
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(2))
    tf = txBox.text_frame
    tf.text = placeholder
    for p in tf.paragraphs:
        p.font.color.rgb = RGBColor(253, 224, 71) # Yellow to highlight it's a placeholder
        p.font.size = Pt(28)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

# --- 10. System Testing ---
add_bullet_slide("9. System Testing", [
    "Unit Testing:",
    "- Tested individual API routes (e.g., verifying weather data returns 200 OK).",
    "Integration Testing:",
    "- Ensured the Flask backend correctly parsed Gemini AI responses before sending to frontend.",
    "User Acceptance Testing (UAT):",
    "- Verified the simulated checkout process correctly updates the SQLite DB and generates the PDF.",
    "UI/UX Testing:",
    "- Ensured responsive layouts across mobile, tablet, and desktop views."
])

# --- 11. Conclusion & Future Enhancements ---
add_bullet_slide("10. Conclusion & Future Enhancements", [
    "Conclusion:",
    "- TravelWorld successfully merges travel inspiration with intelligent management.",
    "- Reduces planning friction via AI and unified data integration.",
    "Future Enhancements:",
    "- Integration of real payment gateways (e.g., actual Stripe API).",
    "- Adding Flight and Hotel booking capabilities.",
    "- Real-time chat support and social sharing for itineraries."
])

# Save presentation
output_path = r'C:\Users\VARSHA P\OneDrive\Desktop\travel1\travelworld-backend-python\backend-python\TravelWorld_Detailed_Presentation.pptx'
prs.save(output_path)
print(f"Presentation generated successfully at {output_path}")
